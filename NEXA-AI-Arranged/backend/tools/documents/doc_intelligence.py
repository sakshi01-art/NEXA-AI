import os
import json
from typing import Dict, List, Optional
from pathlib import Path
import asyncio

class DocumentIntelligence:
    """
    Process, analyze, and query documents:
    PDF, DOCX, XLSX, TXT, MD, and more
    """
    
    def __init__(self, ai_provider, vector_store):
        self.ai = ai_provider
        self.vector_store = vector_store
        self.supported_formats = {
            '.pdf': self._extract_pdf,
            '.docx': self._extract_docx,
            '.xlsx': self._extract_excel,
            '.txt': self._extract_text,
            '.md': self._extract_markdown,
            '.csv': self._extract_csv,
            '.json': self._extract_json,
            '.html': self._extract_html,
            '.pptx': self._extract_powerpoint,
        }
        self.indexed_documents: Dict[str, Dict] = {}
    
    async def process_document(self, file_path: str) -> Dict:
        """Process and index a document"""
        
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                return {'success': False, 'error': 'File not found'}
            
            ext = file_path.suffix.lower()
            
            if ext not in self.supported_formats:
                return {
                    'success': False,
                    'error': f'Unsupported format: {ext}',
                    'supported': list(self.supported_formats.keys())
                }
            
            # Extract content
            extractor = self.supported_formats[ext]
            content = await extractor(str(file_path))
            
            if not content.get('success'):
                return content
            
            # Generate summary
            summary = await self._summarize_document(
                content['text'],
                file_path.name
            )
            
            # Extract key information
            key_info = await self._extract_key_information(content['text'])
            
            # Index in vector store
            doc_id = str(file_path.stem)
            await self.vector_store.store_knowledge(
                topic=f"Document: {file_path.name}",
                information=content['text'][:2000],  # Store first 2000 chars
                source=str(file_path)
            )
            
            # Store metadata
            self.indexed_documents[doc_id] = {
                'path': str(file_path),
                'name': file_path.name,
                'type': ext,
                'summary': summary,
                'key_info': key_info,
                'content': content,
                'word_count': len(content['text'].split()),
                'indexed_at': asyncio.get_event_loop().time()
            }
            
            return {
                'success': True,
                'document_id': doc_id,
                'name': file_path.name,
                'summary': summary,
                'key_info': key_info,
                'word_count': len(content['text'].split()),
                'pages': content.get('pages', 1)
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_pdf(self, file_path: str) -> Dict:
        """Extract text from PDF"""
        
        try:
            import PyPDF2
            
            text_pages = []
            
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                
                for page in reader.pages:
                    text_pages.append(page.extract_text())
            
            full_text = '\n'.join(text_pages)
            
            return {
                'success': True,
                'text': full_text,
                'pages': len(text_pages)
            }
        
        except ImportError:
            # Try pdfplumber as fallback
            try:
                import pdfplumber
                
                text_pages = []
                
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        text_pages.append(page.extract_text() or '')
                
                return {
                    'success': True,
                    'text': '\n'.join(text_pages),
                    'pages': len(text_pages)
                }
            except Exception as e:
                return {'success': False, 'error': str(e)}
    
    async def _extract_docx(self, file_path: str) -> Dict:
        """Extract text from DOCX"""
        
        try:
            from docx import Document
            
            doc = Document(file_path)
            
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join([cell.text for cell in row.cells])
                    table_text.append(row_text)
            
            full_text = '\n'.join(paragraphs + table_text)
            
            return {
                'success': True,
                'text': full_text,
                'paragraphs': len(paragraphs),
                'tables': len(doc.tables)
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_excel(self, file_path: str) -> Dict:
        """Extract data from Excel"""
        
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(file_path, read_only=True)
            
            all_text = []
            sheets_data = {}
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_rows = []
                
                for row in ws.iter_rows(values_only=True):
                    row_text = ' | '.join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        sheet_rows.append(row_text)
                
                sheets_data[sheet_name] = sheet_rows
                all_text.extend(sheet_rows)
            
            return {
                'success': True,
                'text': '\n'.join(all_text),
                'sheets': sheets_data,
                'sheet_count': len(wb.sheetnames)
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_text(self, file_path: str) -> Dict:
        """Extract text from plain text file"""
        
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
            
            return {'success': True, 'text': text}
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_markdown(self, file_path: str) -> Dict:
        """Extract text from Markdown"""
        
        try:
            import markdown
            import re
            
            with open(file_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # Convert to HTML then strip tags
            html = markdown.markdown(md_content)
            clean_text = re.sub('<[^<]+?>', '', html)
            
            return {'success': True, 'text': clean_text, 'markdown': md_content}
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_csv(self, file_path: str) -> Dict:
        """Extract data from CSV"""
        
        try:
            import csv
            
            rows = []
            
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                
                for row in reader:
                    rows.append(' | '.join(row))
            
            return {
                'success': True,
                'text': '\n'.join(rows),
                'row_count': len(rows)
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_json(self, file_path: str) -> Dict:
        """Extract and format JSON"""
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            text = json.dumps(data, indent=2)
            
            return {'success': True, 'text': text, 'data': data}
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_html(self, file_path: str) -> Dict:
        """Extract text from HTML"""
        
        try:
            from bs4 import BeautifulSoup
            
            with open(file_path, 'r', encoding='utf-8') as f:
                html = f.read()
            
            soup = BeautifulSoup(html, 'html.parser')
            
            # Remove script and style elements
            for element in soup(['script', 'style']):
                element.decompose()
            
            text = soup.get_text(separator='\n')
            
            return {'success': True, 'text': text}
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _extract_powerpoint(self, file_path: str) -> Dict:
        """Extract text from PowerPoint"""
        
        try:
            from pptx import Presentation
            
            prs = Presentation(file_path)
            slides_text = []
            
            for i, slide in enumerate(prs.slides):
                slide_content = [f"Slide {i+1}:"]
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text)
                
                slides_text.append('\n'.join(slide_content))
            
            return {
                'success': True,
                'text': '\n\n'.join(slides_text),
                'slide_count': len(prs.slides)
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def _summarize_document(self, text: str, filename: str) -> str:
        """Generate AI summary of document"""
        
        # Truncate long documents
        truncated = text[:5000] if len(text) > 5000 else text
        
        prompt = f"""
        Summarize this document named "{filename}" in Hinglish:
        
        {truncated}
        
        Provide a clear, concise summary covering:
        1. Document type aur purpose
        2. Main points
        3. Key information
        
        Keep it under 150 words.
        """
        
        try:
            response = await self.ai.complete(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.5,
                max_tokens=300
            )
            
            return response.get('content', 'Summary not available')
        except:
            return f"Document: {filename} ({len(text)} characters)"
    
    async def _extract_key_information(self, text: str) -> Dict:
        """Extract structured key information from document"""
        
        prompt = f"""
        Extract key information from this document.
        
        Text: {text[:3000]}
        
        Return JSON:
        {{
            "dates": ["any dates found"],
            "names": ["person names"],
            "organizations": ["company/org names"],
            "numbers": ["important numbers/amounts"],
            "topics": ["main topics covered"],
            "action_items": ["things that need to be done"]
        }}
        """
        
        try:
            response = await self.ai.complete(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.3
            )
            
            return json.loads(response.get('content', '{}'))
        except:
            return {}
    
    async def query_document(
        self,
        document_id: str,
        question: str
    ) -> Dict:
        """Answer questions about a specific document"""
        
        doc = self.indexed_documents.get(document_id)
        
        if not doc:
            return {
                'success': False,
                'error': 'Document not indexed. Please process it first.'
            }
        
        content = doc.get('content', {}).get('text', '')
        
        prompt = f"""
        Document: {doc['name']}
        
        Content:
        {content[:6000]}
        
        Question: {question}
        
        Answer in the same language as the question (Hinglish/Hindi/English).
        Be specific and cite relevant parts of the document.
        If the answer is not in the document, say so clearly.
        """
        
        try:
            response = await self.ai.complete(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.5
            )
            
            return {
                'success': True,
                'answer': response.get('content', ''),
                'document': doc['name']
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def compare_documents(
        self,
        doc_id_1: str,
        doc_id_2: str
    ) -> Dict:
        """Compare two documents"""
        
        doc1 = self.indexed_documents.get(doc_id_1)
        doc2 = self.indexed_documents.get(doc_id_2)
        
        if not doc1 or not doc2:
            return {
                'success': False,
                'error': 'One or both documents not indexed'
            }
        
        prompt = f"""
        Compare these two documents in Hinglish:
        
        Document 1: {doc1['name']}
        {doc1['content']['text'][:2000]}
        
        Document 2: {doc2['name']}
        {doc2['content']['text'][:2000]}
        
        Provide comparison covering:
        1. Main similarities
        2. Key differences
        3. Which is more recent/relevant
        4. Summary recommendation
        """
        
        try:
            response = await self.ai.complete(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.5
            )
            
            return {
                'success': True,
                'comparison': response.get('content', ''),
                'doc1': doc1['name'],
                'doc2': doc2['name']
            }
        
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    async def search_across_documents(self, query: str) -> Dict:
        """Search across all indexed documents"""
        
        if not self.indexed_documents:
            return {
                'success': False,
                'error': 'No documents indexed yet'
            }
        
        # Search vector store
        results = await self.vector_store.query_knowledge(query, limit=5)
        
        relevant_docs = []
        
        for result in results:
            metadata = result.get('metadata', {})
            source = metadata.get('source', '')
            
            if source:
                doc_name = Path(source).stem
                if doc_name in self.indexed_documents:
                    relevant_docs.append({
                        'document': self.indexed_documents[doc_name]['name'],
                        'relevance': result.get('relevance', 0),
                        'excerpt': result.get('information', '')[:200]
                    })
        
        return {
            'success': True,
            'query': query,
            'results': relevant_docs,
            'count': len(relevant_docs)
        }
